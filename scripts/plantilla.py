"""Arma el producto "El tablero de la CEO": plantilla de finanzas personales en Excel (abre en Google Sheets).

Uso: python plantilla.py      → producto/plantilla/el-tablero-de-la-ceo.xlsx

Hojas: Inicio, Presupuesto 50/30/20, Págate primero, Fondo de emergencia, Deudas (bola de nieve y
avalancha), Interés compuesto (con regla del 72) y Metas. Las celdas amarillas son para escribir;
todo lo demás se calcula solo. Las reglas vienen de la guía (producto/guia/reglas.json).
"""
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NEGRO, ORO, CREMA, GRIS = '0B0B0D', 'F5C542', 'FFF6D6', '6B6B6B'
DINERO = '"$"#,##0'
PCT = '0%'
TITULO = Font(name='Arial', size=20, bold=True, color=ORO)
SUB = Font(name='Arial', size=10, italic=True, color=GRIS)
ENCABEZADO = Font(name='Arial', size=10, bold=True, color=ORO)
NORMAL = Font(name='Arial', size=11)
FUERTE = Font(name='Arial', size=11, bold=True)
FONDO_NEGRO = PatternFill('solid', fgColor=NEGRO)
ENTRADA = PatternFill('solid', fgColor=CREMA)
LINEA = Border(bottom=Side(style='thin', color='DDDDDD'))


def hoja(libro, nombre, titulo, subtitulo, anchos):
    h = libro.create_sheet(nombre)
    h.sheet_view.showGridLines = False
    for i, ancho in enumerate(anchos):
        h.column_dimensions[chr(66 + i)].width = ancho
    h.column_dimensions['A'].width = 3
    h['B2'] = titulo
    h['B2'].font = TITULO
    h['B3'] = subtitulo
    h['B3'].font = SUB
    for col in range(2, 2 + len(anchos)):
        h.cell(row=2, column=col).fill = FONDO_NEGRO
    return h


def encabezados(h, fila, textos):
    for i, t in enumerate(textos):
        c = h.cell(row=fila, column=2 + i, value=t)
        c.font, c.fill = ENCABEZADO, FONDO_NEGRO
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    h.row_dimensions[fila].height = 30


def entrada(c, valor=None, formato=None):
    if valor is not None:
        c.value = valor
    c.fill, c.font = ENTRADA, FUERTE
    if formato:
        c.number_format = formato


def dato(h, fila, etiqueta, formula, formato=DINERO, fuerte=False):
    h.cell(row=fila, column=2, value=etiqueta).font = FUERTE if fuerte else NORMAL
    c = h.cell(row=fila, column=3, value=formula)
    c.number_format = formato
    c.font = FUERTE if fuerte else NORMAL
    for col in (2, 3):
        h.cell(row=fila, column=col).border = LINEA
    return c


def inicio(libro):
    h = libro.active
    h.title = 'Inicio'
    h.sheet_view.showGridLines = False
    h.column_dimensions['A'].width = 3
    h.column_dimensions['B'].width = 95
    filas = [
        ('EL TABLERO DE LA CEO', TITULO),
        ('Tu dinero en una sola hoja de control. Por Victoria The CEO (@victoriatheceo).', SUB),
        ('', NORMAL),
        ('Cómo se usa', FUERTE),
        ('1. Escribe solo en las celdas AMARILLAS. Todo lo demás se calcula solo.', NORMAL),
        ('2. Empieza por "Presupuesto": tu ingreso mensual y tus gastos reales.', NORMAL),
        ('3. "Págate primero": decide qué porcentaje te apartas antes de pagar nada (mínimo 10%).', NORMAL),
        ('4. "Fondo de emergencia": junta de 3 a 6 meses de gastos fijos antes de invertir.', NORMAL),
        ('5. "Deudas": anota cada deuda y elige el orden: bola de nieve o avalancha.', NORMAL),
        ('6. "Interés compuesto": mira cuánto crece un ahorro constante con el tiempo.', NORMAL),
        ('7. "Metas": ponle monto y fecha a lo que quieres; la hoja te dice cuánto apartar al mes.', NORMAL),
        ('', NORMAL),
        ('En Google Sheets: Archivo → Importar → Subir este archivo. Funciona igual.', NORMAL),
        ('', NORMAL),
        ('Aviso: herramienta educativa. No es asesoría financiera, legal ni de inversión. Los rendimientos del '
         'interés compuesto son ejemplos hipotéticos, no promesas. Victoria es un personaje creado con IA.', SUB),
    ]
    for i, (texto, fuente) in enumerate(filas, start=2):
        c = h.cell(row=i, column=2, value=texto)
        c.font = fuente
        c.alignment = Alignment(wrap_text=True, vertical='top')
    h['B2'].fill = FONDO_NEGRO
    h.row_dimensions[16].height = 45


def presupuesto(libro):
    h = hoja(libro, 'Presupuesto', 'Presupuesto 50/30/20', 'Lo necesario 50% · Lo que quieres 30% · Ahorro y deudas 20%. Ajústalo a tu vida.', [34, 16, 16, 16, 14])
    h['B5'] = 'Ingreso mensual neto'
    h['B5'].font = FUERTE
    entrada(h['C5'], 20000, DINERO)
    encabezados(h, 7, ['Bolsa', 'Meta', 'Real', 'Diferencia', '% del ingreso'])
    bolsas = [('Necesidades', 0.5), ('Deseos', 0.3), ('Ahorro y deudas', 0.2)]
    for i, (nombre, pct) in enumerate(bolsas, start=8):
        h.cell(row=i, column=2, value=f'{nombre} ({int(pct * 100)}%)').font = FUERTE
        h.cell(row=i, column=3, value=f'=$C$5*{pct}').number_format = DINERO
        h.cell(row=i, column=4, value=f'=SUMIF($C$14:$C$40,"{nombre}",$D$14:$D$40)').number_format = DINERO
        h.cell(row=i, column=5, value=f'=C{i}-D{i}').number_format = DINERO
        h.cell(row=i, column=6, value=f'=IF($C$5=0,0,D{i}/$C$5)').number_format = PCT
    h.conditional_formatting.add('E8:E9', CellIsRule(operator='lessThan', formula=['0'], font=Font(color='C0392B', bold=True)))
    h.conditional_formatting.add('E10', CellIsRule(operator='greaterThan', formula=['0'], font=Font(color='C0392B', bold=True)))
    h['B11'] = 'Te sobra (o falta) en el mes'
    h['B11'].font = FUERTE
    h['C11'] = '=C5-SUM(D8:D10)'
    h['C11'].number_format, h['C11'].font = DINERO, FUERTE
    encabezados(h, 13, ['Gasto', 'Bolsa', 'Monto del mes'])
    gastos = [('Renta o hipoteca', 'Necesidades', 6000), ('Súper y comida', 'Necesidades', 3500),
              ('Luz, agua, gas, internet', 'Necesidades', 1200), ('Transporte', 'Necesidades', 1000),
              ('Salidas y restaurantes', 'Deseos', 2000), ('Suscripciones', 'Deseos', 500),
              ('Ropa y caprichos', 'Deseos', 1500), ('Ahorro (págate primero)', 'Ahorro y deudas', 2000),
              ('Pago extra a deudas', 'Ahorro y deudas', 1000)]
    validacion = DataValidation(type='list', formula1='"Necesidades,Deseos,Ahorro y deudas"', allow_blank=True)
    h.add_data_validation(validacion)
    for i in range(14, 41):
        nombre, bolsa, monto = gastos[i - 14] if i - 14 < len(gastos) else (None, None, None)
        entrada(h.cell(row=i, column=2), nombre)
        entrada(h.cell(row=i, column=3), bolsa)
        entrada(h.cell(row=i, column=4), monto, DINERO)
        validacion.add(h.cell(row=i, column=3))
    h.freeze_panes = 'B5'


def pagate_primero(libro):
    h = hoja(libro, 'Págate primero', 'Págate primero', 'Aparta tu parte el día que cobras, antes de pagar nada. Regla de El hombre más rico de Babilonia.', [26, 18, 18, 18])
    h['B5'], h['B6'] = 'Ingreso mensual', 'Porcentaje que te pagas'
    h['C5'] = '=Presupuesto!C5'
    h['C5'].number_format = DINERO
    entrada(h['C6'], 0.1, PCT)
    dato(h, 7, 'Te apartas cada mes', '=C5*C6', fuerte=True)
    encabezados(h, 9, ['Mes', 'Lo que apartaste', 'Acumulado', '¿Cumpliste?'])
    meses = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
    for i, mes in enumerate(meses, start=10):
        h.cell(row=i, column=2, value=mes).font = NORMAL
        entrada(h.cell(row=i, column=3), None, DINERO)
        h.cell(row=i, column=4, value=f'=SUM($C$10:C{i})').number_format = DINERO
        h.cell(row=i, column=5, value=f'=IF(C{i}="","",IF(C{i}>=$C$7,"Sí","Faltó"))')
    h.conditional_formatting.add('E10:E21', CellIsRule(operator='equal', formula=['"Faltó"'], font=Font(color='C0392B', bold=True)))
    h.conditional_formatting.add('E10:E21', CellIsRule(operator='equal', formula=['"Sí"'], font=Font(color='1E8449', bold=True)))


def fondo(libro):
    h = hoja(libro, 'Fondo de emergencia', 'Fondo de emergencia', 'La CONDUSEF recomienda de 3 a 6 meses de gastos fijos. Primero esto, después invertir.', [30, 18])
    encabezados(h, 5, ['Gasto fijo mensual', 'Monto'])
    fijos = [('Vivienda', 6000), ('Comida', 3500), ('Salud', 800), ('Deudas (pagos mínimos)', 1500), ('Transporte', 1000), ('Servicios', 1200)]
    for i in range(6, 14):
        nombre, monto = fijos[i - 6] if i - 6 < len(fijos) else (None, None)
        entrada(h.cell(row=i, column=2), nombre)
        entrada(h.cell(row=i, column=3), monto, DINERO)
    dato(h, 14, 'Gastos fijos al mes', '=SUM(C6:C13)', fuerte=True)
    h['B16'], h['B17'], h['B18'] = 'Meses que quieres cubrir (3 a 6)', 'Ya tienes ahorrado', 'Puedes apartar al mes'
    entrada(h['C16'], 3, '0')
    entrada(h['C17'], 5000, DINERO)
    entrada(h['C18'], 2000, DINERO)
    dato(h, 20, 'Tu meta', '=C14*C16', fuerte=True)
    dato(h, 21, 'Te falta', '=MAX(0,C20-C17)')
    dato(h, 22, 'Avance', '=IF(C20=0,0,MIN(1,C17/C20))', PCT, fuerte=True)
    dato(h, 23, 'Meses para llegar', '=IF(C21=0,0,IF(C18<=0,"Aparta algo al mes",ROUNDUP(C21/C18,0)))', '0')
    h.conditional_formatting.add('C22', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=ORO))


def deudas(libro):
    h = hoja(libro, 'Deudas', 'Salir de deudas', 'Bola de nieve: primero la más chica (impulso). Avalancha: primero la de tasa más alta (pagas menos).', [24, 15, 13, 15, 15, 15, 15])
    encabezados(h, 5, ['Deuda', 'Saldo', 'Tasa anual', 'Pago mínimo', 'Interés del mes', 'Orden bola de nieve', 'Orden avalancha'])
    ejemplo = [('Tarjeta A', 18000, 0.65, 900), ('Tarjeta B', 6000, 0.45, 300), ('Préstamo personal', 25000, 0.30, 1400)]
    for i in range(6, 16):
        nombre, saldo, tasa, minimo = ejemplo[i - 6] if i - 6 < len(ejemplo) else (None, None, None, None)
        entrada(h.cell(row=i, column=2), nombre)
        entrada(h.cell(row=i, column=3), saldo, DINERO)
        entrada(h.cell(row=i, column=4), tasa, PCT)
        entrada(h.cell(row=i, column=5), minimo, DINERO)
        h.cell(row=i, column=6, value=f'=IF(C{i}="","",C{i}*D{i}/12)').number_format = DINERO
        h.cell(row=i, column=7, value=f'=IF(C{i}="","",COUNTIFS($C$6:$C$15,"<"&C{i})+1)')
        h.cell(row=i, column=8, value=f'=IF(D{i}="","",COUNTIFS($D$6:$D$15,">"&D{i})+1)')
    dato(h, 17, 'Deuda total', '=SUM(C6:C15)', fuerte=True)
    dato(h, 18, 'Pagos mínimos al mes', '=SUM(E6:E15)')
    dato(h, 19, 'Intereses que pagas al mes (aprox.)', '=SUM(F6:F15)', fuerte=True)
    h['B21'] = 'Escribe la tasa anual de cada deuda (está en tu estado de cuenta). El orden 1 es la que atacas primero con todo lo extra; a las demás, solo el mínimo.'
    h['B21'].font = SUB
    h['B21'].alignment = Alignment(wrap_text=True)
    h.merge_cells('B21:H22')


def interes(libro):
    h = hoja(libro, 'Interés compuesto', 'Interés compuesto', 'Ejemplo hipotético: los rendimientos reales cambian y no están garantizados.', [32, 18, 18, 18])
    etiquetas = [('Capital inicial', 0, DINERO), ('Aporte mensual', 1000, DINERO), ('Rendimiento anual esperado', 0.08, '0.0%'), ('Años', 30, '0')]
    for i, (t, v, f) in enumerate(etiquetas, start=5):
        h.cell(row=i, column=2, value=t).font = FUERTE
        entrada(h.cell(row=i, column=3), v, f)
    dato(h, 10, 'Tendrías al final', '=-FV(C7/12,C8*12,C6,C5)', fuerte=True)
    dato(h, 11, 'Tú habrías puesto', '=C5+C6*12*C8')
    dato(h, 12, 'Lo que puso el interés', '=C10-C11', fuerte=True)
    dato(h, 13, 'Regla del 72: años para duplicar', '=IF(C7=0,"",72/(C7*100))', '0.0')
    encabezados(h, 15, ['Año', 'Saldo al final del año', 'Total que pusiste', 'Ganancia acumulada'])
    for anio in range(1, 41):
        fila = 15 + anio
        h.cell(row=fila, column=2, value=f'=IF({anio}>$C$8,"",{anio})')
        h.cell(row=fila, column=3, value=f'=IF({anio}>$C$8,"",-FV($C$7/12,{anio}*12,$C$6,$C$5))').number_format = DINERO
        h.cell(row=fila, column=4, value=f'=IF({anio}>$C$8,"",$C$5+$C$6*12*{anio})').number_format = DINERO
        h.cell(row=fila, column=5, value=f'=IF({anio}>$C$8,"",C{fila}-D{fila})').number_format = DINERO


def metas(libro):
    h = hoja(libro, 'Metas', 'Metas', 'Una meta sin monto y sin fecha es un deseo. Aquí cada una tiene las dos.', [28, 15, 15, 15, 15, 13, 17])
    encabezados(h, 5, ['Meta', 'Cuánto cuesta', 'Fecha', 'Ya tienes', 'Te falta', 'Meses', 'Apartar al mes'])
    ejemplos = [('Fondo de emergencia', 30000, '2027-06-30', 5000), ('Viaje', 25000, '2027-12-15', 0), ('Enganche', 150000, '2029-01-31', 10000)]
    import datetime
    for i in range(6, 16):
        nombre, costo, fecha, tiene = ejemplos[i - 6] if i - 6 < len(ejemplos) else (None, None, None, None)
        entrada(h.cell(row=i, column=2), nombre)
        entrada(h.cell(row=i, column=3), costo, DINERO)
        entrada(h.cell(row=i, column=4), datetime.date.fromisoformat(fecha) if fecha else None, 'dd/mm/yyyy')
        entrada(h.cell(row=i, column=5), tiene, DINERO)
        h.cell(row=i, column=6, value=f'=IF(C{i}="","",MAX(0,C{i}-E{i}))').number_format = DINERO
        h.cell(row=i, column=7, value=f'=IF(D{i}="","",MAX(1,ROUND((D{i}-TODAY())/30.4,0)))')
        h.cell(row=i, column=8, value=f'=IF(OR(C{i}="",D{i}=""),"",F{i}/G{i})').number_format = DINERO
    h['B17'] = 'Si "Apartar al mes" no te alcanza, mueve la fecha o baja el monto: es mejor una meta real que una perfecta.'
    h['B17'].font = SUB


def main():
    libro = Workbook()
    inicio(libro)
    presupuesto(libro)
    pagate_primero(libro)
    fondo(libro)
    deudas(libro)
    interes(libro)
    metas(libro)
    salida = os.path.join(RAIZ, 'producto', 'plantilla', 'el-tablero-de-la-ceo.xlsx')
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    libro.save(salida)
    print('Plantilla lista:', salida)


if __name__ == '__main__':
    main()
